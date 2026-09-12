"""RAG 离线入库链（对齐企业级 RAG 文档第二节：解析 → 清洗 → 分块 → 元数据 → 嵌入）

上传文档落盘 → 按类型加载器解析 → 清洗 → 递归分块（父块 + 子块）→ 治理元数据 → 向量化入库

支持格式：
- 文本类：PDF、Word(.docx)、PowerPoint(.pptx)、Excel(.xlsx)、TXT、Markdown、HTML
- 图片类：PNG/JPG（OCR，需 pytesseract + Tesseract）
- 扫描件 PDF：无文字层时自动识别并转 OCR（需 PyMuPDF 渲染），引擎缺失则给出安装指引

保留结构：Word 按 Heading、Markdown 按 # 标题、HTML 按 h1~h3、PPT 按幻灯片（标题作章节）、
Excel 按工作表切段；页码/章节随分块落库，回答引用可定位到「文件 > 章节 > 页码」。
"""
import os
import re
import time

from app.config import settings
from app.services import governance, ocr, rag
from app.services import parents as parent_store

SUPPORTED_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".html", ".htm", ".png", ".jpg", ".jpeg"}
IMAGE_EXT = {".png", ".jpg", ".jpeg"}

# 返回结构：[(文本, 页码, 章节名)]
PageBlock = tuple[str, int, str]


# ---------------- 1. 类型加载器 ----------------

class _HTMLText:
    """HTML 正文提取：跳过 script/style，h1~h3 作为章节"""

    _SKIP = {"script", "style", "noscript"}
    _HEAD = {"h1", "h2", "h3"}

    def __init__(self):
        from html.parser import HTMLParser

        outer = self

        class _P(HTMLParser):
            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.skip = 0
                self.heading = False
                self.section = ""
                self.buf: list[str] = []
                self.blocks: list[tuple[str, str]] = []

            def handle_starttag(self, tag, attrs):
                if tag in outer._SKIP:
                    self.skip += 1
                elif tag in outer._HEAD:
                    self._flush()
                    self.heading = True

            def handle_endtag(self, tag):
                if tag in outer._SKIP and self.skip:
                    self.skip -= 1
                elif tag in outer._HEAD:
                    self.heading = False

            def handle_data(self, data):
                if self.skip:
                    return
                text = data.strip()
                if not text:
                    return
                if self.heading:
                    self.section = text
                else:
                    self.buf.append(text)

            def _flush(self):
                if self.buf:
                    self.blocks.append((self.section, "\n".join(self.buf)))
                    self.buf = []

            def finish(self):
                self._flush()
                return self.blocks

        self._parser = _P()

    def parse(self, raw: str) -> list[tuple[str, str]]:
        self._parser.feed(raw)
        return self._parser.finish()


def _split_by_heading(text: str) -> list[PageBlock]:
    """按 Markdown 风格标题（#）切段；无标题则整篇一段（txt/md 通用）"""
    blocks: list[PageBlock] = []
    section = ""
    buf: list[str] = []
    for line in text.splitlines():
        stripped = line.lstrip()
        if stripped.startswith("#"):
            if buf:
                blocks.append(("\n".join(buf), 1, section))
                buf = []
            section = stripped.lstrip("#").strip()
            continue
        buf.append(line)
    if buf:
        blocks.append(("\n".join(buf), 1, section))
    return blocks or [(text, 1, "")]


def load_pages(path: str, ext: str) -> list[PageBlock]:
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages: list[PageBlock] = [((page.extract_text() or ""), i + 1, f"第{i + 1}页")
                                 for i, page in enumerate(reader.pages)]
        # 文字层过少 → 判定为扫描件，尝试 OCR（引擎缺失则抛出带指引的错误）
        letters = sum(len(t.strip()) for t, _, _ in pages)
        if letters < 20 * max(1, len(pages)):
            st = ocr.engine_status()
            if st["available"] and st["pymupdf"]:
                return ocr.pdf_scan_to_text(path)
            raise ValueError("该 PDF 无文字层（疑似扫描件），无法直接解析。" + (st["hint"] or ""))
        return pages

    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        blocks: list[PageBlock] = []
        section = ""
        buf: list[str] = []

        def flush() -> None:
            if buf:
                blocks.append(("\n".join(buf), 1, section))
                buf.clear()

        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue
            style = (getattr(p.style, "name", "") or "")
            if style.startswith("Heading") or style.startswith("标题"):
                flush()
                section = text
                continue
            buf.append(text)
        flush()
        return blocks or [("", 1, "")]

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        pages: list[PageBlock] = []
        for i, slide in enumerate(prs.slides):
            parts: list[str] = []
            title = ""
            try:
                if slide.shapes.title and slide.shapes.title.text:
                    title = slide.shapes.title.text.strip()
            except Exception:
                title = ""
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False):
                    try:
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells if c.text.strip()]
                            if cells:
                                parts.append(" | ".join(cells))
                    except Exception:
                        pass
            pages.append(("\n".join(parts), i + 1, title or f"幻灯片{i + 1}"))
        return pages

    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        pages: list[PageBlock] = []
        for si, ws in enumerate(wb.worksheets):
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
            pages.append(("\n".join(rows), si + 1, ws.title or f"工作表{si + 1}"))
        return pages

    if ext in (".html", ".htm"):
        with open(path, encoding="utf-8", errors="ignore") as f:
            raw = f.read()
        blocks = _HTMLText().parse(raw)
        if not blocks:
            return [("", 1, "")]
        return [(text, 1, section) for section, text in blocks]

    if ext in IMAGE_EXT:
        return [(ocr.image_to_text(path), 1, "图片OCR")]

    # txt / md：md 按标题切，txt 整篇一段
    with open(path, encoding="utf-8", errors="ignore") as f:
        content = f.read()
    return _split_by_heading(content) if ext == ".md" else [(content, 1, "")]


# ---------------- 2. 清洗 ----------------

def clean_text(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)               # 压缩水平空白
    text = re.sub(r"\n{3,}", "\n\n", text)            # 压缩连续空行
    text = re.sub(r"-{3,}|={3,}|\*{3,}", "", text)    # 去分隔线
    # 连续重复行折叠（扫描件/页眉页脚造成的重复文本）
    lines, out, prev, repeat = text.splitlines(), [], None, 0
    for line in lines:
        if line.strip() and line == prev:
            repeat += 1
            if repeat >= 1:      # 相同行只保留一次
                continue
        else:
            repeat = 0
        out.append(line)
        prev = line
    return "\n".join(out).strip()


# ---------------- 3. 递归语义分块 ----------------

_SEPARATORS = ["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]


def recursive_split(text: str, size: int, overlap: int) -> list[str]:
    """递归字符分块：优先按大分隔符切，超长再细分；相邻块保留 overlap 重叠"""
    if len(text) <= size:
        return [text] if text.strip() else []

    for sep in _SEPARATORS:
        if sep and sep in text:
            parts = text.split(sep)
            break
    else:
        parts = [text[i:i + size] for i in range(0, len(text), size - overlap)]
        return [p for p in parts if p.strip()]

    chunks: list[str] = []
    buf = ""
    for part in parts:
        piece = (sep if sep and not buf.endswith(sep) else "") + part if buf else part
        if not piece.strip():
            continue
        if len(buf) + len(piece) + 1 <= size:
            buf = f"{buf}{sep}{piece}" if buf else piece
        else:
            if buf.strip():
                chunks.append(buf)
            # 长度仍超限的 piece 递归细分
            if len(piece) > size:
                chunks.extend(recursive_split(piece, size, overlap))
                buf = ""
            else:
                tail = buf[-overlap:] if buf and overlap > 0 else ""
                buf = tail + piece
    if buf.strip():
        chunks.append(buf)
    return chunks


# ---------------- 4. 完整入库链 ----------------

def ingest_file(path: str, filename: str, category: str, *, ctx: dict | None = None,
                security_level: str | None = None, dept_id: str | None = None,
                review_status: str | None = None, effective_ts: float = 0.0,
                expire_ts: float = 0.0, doc_version: int = 1, title: str = "") -> dict:
    """解析→清洗→分块（父块+子块）→治理元数据→向量化入库。返回统计信息。"""
    ctx = ctx or governance.access_context()
    ext = os.path.splitext(filename)[1].lower()
    blocks = load_pages(path, ext)

    all_chunks: list[dict] = []
    parent_items: dict[str, str] = {}
    parent_size = settings.RAG_CHUNK_SIZE * max(1, settings.RAG_PARENT_MULTIPLIER)
    for text, page, section in blocks:
        text = clean_text(text)
        if not text:
            continue
        # 父块：更大粒度，仅存文本供生成时补全上下文（不参与检索，小块检索大块生成）
        parents = recursive_split(text, parent_size, settings.RAG_CHUNK_OVERLAP) if settings.RAG_PARENT_ENABLED else []
        for idx, chunk in enumerate(recursive_split(text, settings.RAG_CHUNK_SIZE, settings.RAG_CHUNK_OVERLAP)):
            parent_id = ""
            if parents:
                p_idx = next((i for i, pt in enumerate(parents) if chunk[:40] in pt), 0)
                parent_id = f"{filename}#{page}#{p_idx}"
                parent_items.setdefault(parent_id, parents[p_idx])
            all_chunks.append({
                "text": chunk,
                "metadata": governance.chunk_meta(
                    source=filename, category=category, page=page, chunk_index=idx, ctx=ctx,
                    security_level=security_level, dept_id=dept_id, review_status=review_status,
                    effective_ts=effective_ts, expire_ts=expire_ts, doc_version=doc_version,
                    title=title or filename, section=section, parent_id=parent_id,
                ),
            })

    n = rag.add_chunks(all_chunks)
    parents_saved = parent_store.put_many(parent_items) if parent_items else 0
    return {"filename": filename, "pages": len(blocks), "chunks": n, "parents": parents_saved,
            "security_level": governance.normalize_level(security_level),
            "review_status": governance.normalize_review(review_status)}


def save_upload(filename: str, content: bytes) -> str:
    """原始文件落盘"""
    os.makedirs(settings.RAG_UPLOAD_DIR, exist_ok=True)
    safe = f"{int(time.time() * 1000)}_{re.sub(r'[\\\\/:*?\"<>|]', '_', filename)}"
    path = os.path.join(settings.RAG_UPLOAD_DIR, safe)
    with open(path, "wb") as f:
        f.write(content)
    return path
