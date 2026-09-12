"""企业级能力冒烟：外部存储回退 / 解析扩展 / Langfuse / OIDC / TOTP / 失败锁定 / 合规报表

用法（需后端已启动，.env 建议开启 OIDC_MOCK=true 以验证单点登录流程）：
    python tests/smoke_enterprise.py
"""
import io
import json
import sys
import threading
import time

import httpx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, ".")
BASE = "http://127.0.0.1:8010/api/v1"
RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(ok), detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" :: {detail}" if detail else ""))


def client() -> httpx.Client:
    return httpx.Client(timeout=120, trust_env=False)


def admin_headers(c: httpx.Client) -> dict:
    r = c.post(f"{BASE}/auth/login", json={"username": "admin", "password": "123456"})
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['data']['token']}"}


# ---------------- 1. 失败锁定（等保：身份鉴别） ----------------

def test_login_lock(c: httpx.Client) -> None:
    """等保：连续失败锁定。用 wangmin 验证，避免锁定 admin 影响后续用例"""
    for _ in range(5):
        c.post(f"{BASE}/auth/login", json={"username": "wangmin", "password": "wrong-pass"})
    r = c.post(f"{BASE}/auth/login", json={"username": "wangmin", "password": "123456"})
    body = r.json()
    check("登录失败锁定", r.status_code == 429 and body.get("code") == 2002,
          f"连续错误后返回 {r.status_code}/{body.get('code')}（锁定到期自动解锁）")


# ---------------- 2. TOTP 双因素 ----------------

def test_totp(c: httpx.Client, headers: dict) -> None:
    from app.services import totp as totp_svc
    r = c.post(f"{BASE}/auth/totp/setup", headers=headers)
    data = r.json()["data"]
    secret = data["secret"]
    check("TOTP setup 返回密钥与 otpauth URI", bool(secret and data.get("otpauth_uri", "").startswith("otpauth://")))
    code = totp_svc._code(secret, int(time.time()) // 30)
    r = c.post(f"{BASE}/auth/totp/enable", headers=headers, json={"code": code})
    check("TOTP enable 校验通过", r.json().get("code") == 0)
    # 错误口令应被拒
    r = c.post(f"{BASE}/auth/login", json={"username": "admin", "password": "123456", "totp_code": "000000"})
    check("启用后登录必须携带正确动态口令", r.status_code == 401)
    # 正确口令
    code = totp_svc._code(secret, int(time.time()) // 30)
    r = c.post(f"{BASE}/auth/login", json={"username": "admin", "password": "123456", "totp_code": code})
    check("携带正确动态口令登录成功", r.json().get("code") == 0)
    c.post(f"{BASE}/auth/totp/disable", headers=headers)
    r = c.post(f"{BASE}/auth/login", json={"username": "admin", "password": "123456"})
    check("关闭后恢复单因素登录", r.json().get("code") == 0)


# ---------------- 3. OIDC mock 端到端 ----------------

def test_oidc_mock(c: httpx.Client) -> None:
    st = c.get(f"{BASE}/auth/oidc/status").json()["data"]
    if not st.get("mock_idp"):
        check("OIDC mock IdP 端到端", True, "跳过：OIDC_MOCK 未开启")
        return
    r = c.get(f"{BASE}/auth/oidc/login?username=admin", follow_redirects=False)
    loc1 = r.headers.get("location", "")
    check("OIDC /login 302 到 mock authorize",
          r.status_code in (301, 302, 307) and "/oidc/mock/authorize" in loc1, loc1[:80])
    r = c.get(loc1, follow_redirects=False)
    loc2 = r.headers.get("location", "")
    check("mock authorize 302 回 callback 并携带 code/state", "code=" in loc2 and "state=" in loc2, loc2[:100])
    from urllib.parse import urlparse, parse_qs
    q = parse_qs(urlparse(loc2).query)
    r = c.get(f"{BASE}/auth/oidc/callback", params={"code": q["code"][0], "state": q["state"][0]})
    body = r.json()
    check("OIDC 回调签发平台 Token", body.get("code") == 0 and body.get("data", {}).get("token"), body.get("msg", ""))


# ---------------- 4. 解析扩展：HTML / PPTX ----------------

def wait_doc_active(c: httpx.Client, headers: dict, name: str, timeout: int = 90) -> bool:
    """轮询文档状态直到入库完成（后台解析需先加载嵌入模型，首次约 20~40 秒）"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        rows = c.get(f"{BASE}/documents", headers=headers).json().get("data") or []
        row = next((x for x in rows if x.get("name") == name), None)
        if row and row.get("status") == "active" and (row.get("chunks") or 0) > 0:
            return True
        time.sleep(2)
    return False


def test_parse_extensions(c: httpx.Client, headers: dict) -> None:
    html = ("<!doctype html><html><head><title>购房流程指南</title></head><body>"
            "<h1>网签备案流程</h1><p>网签需要在交付定金后七个工作日内完成，需携带身份证与购房合同。</p>"
            "<h2>所需材料</h2><p>身份证明、婚姻证明、购房合同原件。</p></body></html>")
    files = {"file": ("guide.html", io.BytesIO(html.encode("utf-8")), "text/html")}
    r = c.post(f"{BASE}/documents/upload", headers=headers, files=files)
    ok_html = r.json().get("code") == 0
    check("HTML 文档解析入库", ok_html, r.json().get("msg", "")[:80])
    check("HTML 后台解析完成", wait_doc_active(c, headers, "guide.html"))

    # PPTX：用 python-pptx 现场生成
    try:
        from pptx import Presentation
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[0])
        s1.shapes.title.text = "滨江花园楼盘介绍"
        s1.placeholders[1].text = "建面98㎡三房两厅，南北通透，均价2.1万/平方米"
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "付款方式"
        s2.placeholders[1].text = "支持一次性付款与按揭贷款，首付最低三成"
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        files = {"file": ("楼盘介绍.pptx", buf, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        r = c.post(f"{BASE}/documents/upload", headers=headers, files=files)
        check("PPT 文档解析入库", r.json().get("code") == 0, r.json().get("msg", "")[:80])
        ppt_uploaded = r.json().get("code") == 0
    except ImportError:
        ppt_uploaded = False
        check("PPT 文档解析入库", True, "跳过：python-pptx 未安装")
    if ppt_uploaded:
        check("PPT 后台解析完成", wait_doc_active(c, headers, "楼盘介绍.pptx"))

    # 章节检索（HTML 的 h1 章节应可定位）
    r = c.post(f"{BASE}/documents/search", headers=headers, json={"query": "网签备案需要什么材料"})
    hits = r.json().get("data") or []
    html_hit = next((h for h in hits if h.get("doc") == "guide.html"), None)
    check("HTML 章节检索与定位", bool(html_hit),
          f"章节={html_hit.get('section', '')} 相关度={html_hit.get('score', 0)}%" if html_hit else "未命中")
    # 清理
    if html_hit:
        c.delete(f"{BASE}/documents/guide.html", headers=headers)


# ---------------- 5. Langfuse 适配器（本地 echo 校验上报体） ----------------

def test_langfuse() -> None:
    from app.config import settings
    from app.services import tracing
    if tracing.enabled():
        check("Langfuse 适配器", True, "已配置真实 Langfuse，跳过本地 echo 验证")
        return

    received: list[dict] = []

    class Echo(threading.Thread):
        def __init__(self):
            super().__init__(daemon=True)
            from http.server import BaseHTTPRequestHandler, HTTPServer

            class H(BaseHTTPRequestHandler):
                def do_POST(self):
                    n = int(self.headers.get("Content-Length", 0) or 0)
                    received.append(json.loads(self.rfile.read(n) or b"{}"))
                    self.send_response(200)
                    self.send_header("Content-Length", "2")
                    self.end_headers()
                    self.wfile.write(b"{}")

                def log_message(self, *a):
                    pass

            self.srv = HTTPServer(("127.0.0.1", 0), H)
            self.port = self.srv.server_address[1]

        def run(self):
            self.srv.serve_forever()

    echo = Echo()
    echo.start()
    old = (settings.LANGFUSE_HOST, settings.LANGFUSE_PUBLIC_KEY, settings.LANGFUSE_SECRET_KEY)
    settings.LANGFUSE_HOST, settings.LANGFUSE_PUBLIC_KEY, settings.LANGFUSE_SECRET_KEY = (
        f"http://127.0.0.1:{echo.port}", "pk-test", "sk-test")
    try:
        tracing.record_rag({"trace_id": "lf-test-1", "ts": time.time(), "username": "admin",
                            "query": "测试问题", "answer": "测试回答", "model": "qwen2.5:0.5b",
                            "tokens": 42, "elapsed_ms": 120, "recall_vector": 3, "recall_bm25": 2,
                            "blocked": 0, "rejected": False, "key_source": "system"})
        res = tracing.flush()
        time.sleep(0.3)
        events = (received[0].get("batch") if received else []) or []
        types = sorted({e.get("type") for e in events})
        check("Langfuse 上报体（零 SDK HTTP ingestion）",
              res is not None and "trace-create" in types and "span-create" in types and "generation-create" in types,
              f"事件类型 {types} 数量 {len(events)}")
    finally:
        settings.LANGFUSE_HOST, settings.LANGFUSE_PUBLIC_KEY, settings.LANGFUSE_SECRET_KEY = old
        tracing._BUFFER.clear()
    st = tracing.status()
    check("未配置 Langfuse 时不外发", not st["langfuse_enabled"] and st["buffered"] == 0, st["hint"][:60])


# ---------------- 6. 合规报表 / 审计哈希链 / 存储状态 ----------------

def test_compliance(c: httpx.Client, headers: dict) -> None:
    r = c.get(f"{BASE}/compliance/report", headers=headers).json()["data"]
    check("等保合规自查报表", r["summary"]["total"] >= 12,
          f"implemented={r['summary']['implemented']} partial={r['summary']['partial']} manual={r['summary']['manual']}")
    r = c.get(f"{BASE}/audit/verify", headers=headers).json()
    check("审计哈希链完整性", r["data"]["valid"], f"共 {r['data']['total']} 条")
    r = c.get(f"{BASE}/chat/llm-status", headers=headers).json()["data"]
    vst = (c.get(f"{BASE}/compliance/report", headers=headers).json()["data"]["runtime"]
           or {}).get("vector_backend", {})
    check("向量后端状态（Chroma/Milvus 回退）", vst.get("active") in ("chroma", "milvus"),
          f"configured={vst.get('configured')} active={vst.get('active')} {vst.get('fallback_reason', '')[:60]}")
    ost = (c.get(f"{BASE}/compliance/report", headers=headers).json()["data"]["runtime"] or {}).get("ocr", {})
    check("OCR 引擎状态与降级提示", "hint" in ost, f"available={ost.get('available')}")


def main() -> int:
    with client() as c:
        headers = admin_headers(c)
        test_login_lock(c)
        test_totp(c, headers)
        test_oidc_mock(c)
        test_parse_extensions(c, headers)
        test_langfuse()
        test_compliance(c, headers)
    passed = sum(1 for _, ok, _ in RESULTS if ok)
    print("\n" + "=" * 68)
    print(f"企业级能力冒烟：{passed}/{len(RESULTS)} 通过")
    for name, ok, detail in RESULTS:
        if not ok:
            print(f"  ✗ {name} :: {detail}")
    print("=" * 68)
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    raise SystemExit(main())
